import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from shockbridge_state_risk.data.ecb_events import (
    EA_EMPD_MINIMUM_COLUMNS,
    EA_MPD_MINIMUM_COLUMNS,
    EA_MPD_SHEETS,
    WorkbookValidationError,
    audit_ea_empd,
    audit_ea_mpd,
)


class EcbEventAuditTests(unittest.TestCase):
    def _ea_mpd_fixture(self, path: Path, duplicate: bool = False) -> None:
        workbook = Workbook()
        notes = workbook.active
        notes.title = "Notes"
        notes.append(["* Last update: 4 November 2025"])
        for name in EA_MPD_SHEETS:
            sheet = workbook.create_sheet(name)
            sheet.append(list(EA_MPD_MINIMUM_COLUMNS))
            sheet.append([datetime(2020, 1, 1)] + [1.0] * (len(EA_MPD_MINIMUM_COLUMNS) - 1))
            second_date = datetime(2020, 1, 1) if duplicate else "02/02/2020"
            sheet.append([second_date] + [None] * (len(EA_MPD_MINIMUM_COLUMNS) - 1))
        workbook.save(path)

    def _ea_empd_fixture(self, path: Path, bad_type: bool = False) -> None:
        workbook = Workbook()
        notes = workbook.active
        notes.title = "Notes"
        sheet = workbook.create_sheet("EA-EMPD")
        sheet.append(list(EA_EMPD_MINIMUM_COLUMNS))
        event_type = "UNKNOWN" if bad_type else "GC_PR"
        sheet.append(
            [datetime(2020, 1, 1, 13, 45), event_type, "Speaker", "Title"]
            + [0.0] * (len(EA_EMPD_MINIMUM_COLUMNS) - 4)
        )
        workbook.save(path)

    def test_original_workbook_audit_handles_mixed_dates(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ea_mpd.xlsx"
            self._ea_mpd_fixture(path)
            audit = audit_ea_mpd(path)
        self.assertEqual(audit.workbook_last_update, "4 November 2025")
        self.assertEqual(audit.sheets["Press Release Window"].rows, 2)
        self.assertEqual(audit.sheets["Press Release Window"].missing_by_column["EURUSD"], 1)

    def test_original_duplicate_date_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ea_mpd.xlsx"
            self._ea_mpd_fixture(path, duplicate=True)
            with self.assertRaises(WorkbookValidationError):
                audit_ea_mpd(path)

    def test_extended_workbook_uses_composite_event_key(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ea_empd.xlsx"
            self._ea_empd_fixture(path)
            audit = audit_ea_empd(path)
        self.assertEqual(audit.rows, 1)
        self.assertEqual(audit.event_type_counts, {"GC_PR": 1})
        self.assertEqual(audit.duplicate_composite_keys, 0)

    def test_extended_unknown_event_type_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ea_empd.xlsx"
            self._ea_empd_fixture(path, bad_type=True)
            with self.assertRaises(WorkbookValidationError):
                audit_ea_empd(path)


if __name__ == "__main__":
    unittest.main()
