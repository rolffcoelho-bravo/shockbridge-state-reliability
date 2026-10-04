"""Schema and feasibility audits for the ECB EA-MPD and EA-EMPD workbooks."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook  # type: ignore[import-untyped]

from shockbridge_state_risk.data.download import sha256_file


class WorkbookValidationError(ValueError):
    """Raised when an event-study workbook violates a structural contract."""


EA_MPD_SHEETS = (
    "Press Release Window",
    "Press Conference Window",
    "Monetary Event Window",
)

EA_MPD_MINIMUM_COLUMNS = (
    "date",
    "OIS_1M",
    "DE2Y",
    "DE10Y",
    "IT10Y",
    "STOXX50",
    "SX7E",
    "EURUSD",
)

EA_EMPD_MINIMUM_COLUMNS = (
    "Date_time",
    "Event_type",
    "Speaker",
    "Title",
    "ECB_database",
    "Non_regular_trading_day",
    "Outside_regular_trading_hours",
    "Days_until_next_GC",
    "OIS_1M",
    "DE2Y",
    "DE10Y",
    "IT10Y",
    "STOXX50E",
    "SX7E",
    "EURUSD",
)

EXPECTED_EVENT_TYPES = {"GC_ME", "GC_PR", "GC_PC", "EB", "P"}


@dataclass(frozen=True)
class SheetAudit:
    rows: int
    columns: int
    first_event: str
    last_event: str
    duplicate_keys: int
    missing_by_column: dict[str, int]


@dataclass(frozen=True)
class EAMPDAudit:
    sha256: str
    bytes: int
    workbook_last_update: Optional[str]
    sheets: dict[str, SheetAudit]
    timing_break: str = "2022-07"
    evidence_status: str = "AUDITED_SOURCE_DATA"


@dataclass(frozen=True)
class EAEMPDAudit:
    sha256: str
    bytes: int
    rows: int
    columns: int
    first_event: str
    last_event: str
    duplicate_composite_keys: int
    event_type_counts: dict[str, int]
    missing_by_column: dict[str, int]
    timing_break: str = "2022-07"
    evidence_status: str = "AUDITED_SOURCE_DATA"


def _headers(row: tuple[Any, ...]) -> list[str]:
    values: list[str] = []
    for value in row:
        if value is None:
            break
        values.append(str(value).strip())
    if len(values) != len(set(values)):
        raise WorkbookValidationError("Duplicate column names detected.")
    return values


def _parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            continue
    raise WorkbookValidationError(f"Unsupported event date: {value!r}")


def _notes_update(notes: list[str]) -> Optional[str]:
    prefix = "last update:"
    for line in notes:
        if prefix in line.lower():
            return line.split(":", 1)[1].strip()
    return None


def _read_notes(workbook: Any) -> list[str]:
    if "Notes" not in workbook.sheetnames:
        return []
    return [
        str(cell).strip()
        for row in workbook["Notes"].iter_rows(values_only=True)
        for cell in row
        if cell is not None and str(cell).strip()
    ]


def audit_ea_mpd(path: Path) -> EAMPDAudit:
    """Audit the scheduled Governing Council event-study workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    missing_sheets = sorted(set(EA_MPD_SHEETS) - set(workbook.sheetnames))
    if missing_sheets:
        raise WorkbookValidationError(f"Missing sheets: {missing_sheets}")

    audits: dict[str, SheetAudit] = {}
    for sheet_name in EA_MPD_SHEETS:
        worksheet = workbook[sheet_name]
        iterator = worksheet.iter_rows(values_only=True)
        headers = _headers(next(iterator))
        missing_columns = sorted(set(EA_MPD_MINIMUM_COLUMNS) - set(headers))
        if missing_columns:
            raise WorkbookValidationError(f"{sheet_name} missing columns: {missing_columns}")
        rows = [row[: len(headers)] for row in iterator if row[0] is not None]
        dates = [_parse_date(row[0]) for row in rows]
        if dates != sorted(dates):
            raise WorkbookValidationError(f"{sheet_name} dates are not sorted.")
        duplicate_count = len(dates) - len(set(dates))
        if duplicate_count:
            raise WorkbookValidationError(f"{sheet_name} contains duplicate dates.")
        missing = {
            header: sum(row[index] is None for row in rows) for index, header in enumerate(headers)
        }
        audits[sheet_name] = SheetAudit(
            rows=len(rows),
            columns=len(headers),
            first_event=min(dates).isoformat(),
            last_event=max(dates).isoformat(),
            duplicate_keys=duplicate_count,
            missing_by_column={key: value for key, value in missing.items() if value},
        )

    notes = _read_notes(workbook)
    return EAMPDAudit(
        sha256=sha256_file(path),
        bytes=path.stat().st_size,
        workbook_last_update=_notes_update(notes),
        sheets=audits,
    )


def audit_ea_empd(path: Path) -> EAEMPDAudit:
    """Audit the extended event universe without treating timestamp reuse as duplication."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "EA-EMPD" not in workbook.sheetnames:
        raise WorkbookValidationError("Missing EA-EMPD sheet.")
    worksheet = workbook["EA-EMPD"]
    iterator = worksheet.iter_rows(values_only=True)
    headers = _headers(next(iterator))
    missing_columns = sorted(set(EA_EMPD_MINIMUM_COLUMNS) - set(headers))
    if missing_columns:
        raise WorkbookValidationError(f"EA-EMPD missing columns: {missing_columns}")
    rows = [row[: len(headers)] for row in iterator if row[0] is not None]
    index = {header: offset for offset, header in enumerate(headers)}
    timestamps = [row[index["Date_time"]] for row in rows]
    if not all(isinstance(value, datetime) for value in timestamps):
        raise WorkbookValidationError("EA-EMPD Date_time must contain Excel datetimes.")
    if timestamps != sorted(timestamps):
        raise WorkbookValidationError("EA-EMPD timestamps are not sorted.")

    event_types = Counter(str(row[index["Event_type"]]) for row in rows)
    unexpected_types = set(event_types) - EXPECTED_EVENT_TYPES
    if unexpected_types:
        raise WorkbookValidationError(f"Unexpected event types: {sorted(unexpected_types)}")
    keys = [
        (
            row[index["Date_time"]],
            row[index["Event_type"]],
            row[index["Speaker"]],
            row[index["Title"]],
        )
        for row in rows
    ]
    duplicate_count = len(keys) - len(set(keys))
    if duplicate_count:
        raise WorkbookValidationError("EA-EMPD contains duplicate composite event keys.")
    missing = {
        header: sum(row[offset] is None for row in rows) for offset, header in enumerate(headers)
    }
    return EAEMPDAudit(
        sha256=sha256_file(path),
        bytes=path.stat().st_size,
        rows=len(rows),
        columns=len(headers),
        first_event=min(timestamps).isoformat(),
        last_event=max(timestamps).isoformat(),
        duplicate_composite_keys=duplicate_count,
        event_type_counts=dict(sorted(event_types.items())),
        missing_by_column={key: value for key, value in missing.items() if value},
    )
